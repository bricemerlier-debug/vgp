"""Create the editable, one-slide monthly report from the supplied deck's background.
The background asset is rendered from slide 1 after removing its original figures.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

root=Path(__file__).resolve().parents[1]
prs=Presentation()
prs.slide_width=Inches(13.333333)
prs.slide_height=Inches(7.5)
slide=prs.slides.add_slide(prs.slide_layouts[6])
slide.shapes.add_picture(str(root/'assets/container-report-background.png'),0,0,width=prs.slide_width,height=prs.slide_height)

def box(text,left,top,width,height,size,bold=False,color=(0,0,0),font='Arial',align=PP_ALIGN.LEFT):
 shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height))
 frame=shape.text_frame
 frame.clear();frame.word_wrap=False
 frame.margin_left=frame.margin_right=Inches(.02)
 frame.margin_top=frame.margin_bottom=0
 paragraph=frame.paragraphs[0];paragraph.alignment=align
 run=paragraph.add_run();run.text=text
 run.font.name=font;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(*color)
 return shape

box('Bilan mois de {{MONTH}}',.3,.85,8.1,.75,39,True,(255,189,41),'Arial Narrow')
labels=['Pli AMIENS','Cartouches PLT2','Cassettes R2 - BLEU','Cassettes R2 - JAUNE','Cassettes R2 - VERTE','Carcasses - PLAT','Carcasses - BACHE','GDF','PC','Chariots','Tracteurs','Divers']
rows=13
left,top,width,height=.46,2.29,5.71,3.12
table=slide.shapes.add_table(rows,3,Inches(left),Inches(top),Inches(width),Inches(height)).table
for i,w in enumerate((1.75,2.75,1.21)):table.columns[i].width=Inches(w)
table.rows[0].height=Inches(.45)
for r in range(1,rows):table.rows[r].height=Inches((height-.45)/12)
table.cell(1,2).merge(table.cell(12,2))
for r in range(rows):
 for c in range(3):
  if c==2 and r>1:continue
  cell=table.cell(r,c);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(*(0,0,0) if r==0 else (255,255,255))
  tcPr=cell._tc.get_or_add_tcPr()
  for side in ('lnL','lnR','lnT','lnB'):
   edge=OxmlElement('a:'+side);edge.set('w','6350')
   fill=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr');color.set('val','222222');fill.append(color);edge.append(fill);tcPr.append(edge)
  cell.margin_left=cell.margin_right=Inches(.02)
  cell.margin_top=cell.margin_bottom=0
  cell.vertical_anchor=MSO_ANCHOR.MIDDLE
  cell.text=''
  p=cell.text_frame.paragraphs[0];p.alignment=PP_ALIGN.CENTER
  run=p.add_run()
  if r==0:run.text=('Types de contenants','Nombre d’interventions','Total')[c]
  elif c==0:run.text=labels[r-1]
  elif c==1:run.text='{{COUNT_%d}}'%(r-1)
  elif c==2:run.text='{{TOTAL}}'
  run.font.name='Arial';run.font.size=Pt(47 if r==1 and c==2 else 8.8 if r else 9.5);run.font.bold=r==0
  run.font.color.rgb=RGBColor(*(255,255,255) if r==0 else (0,0,0))
prs.save(root/'assets/container-report-template.pptx')
